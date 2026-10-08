"""Canonical ingestion contract shared with the manual frontend (fourwall_contract.json)."""
import hashlib
import json
import math
import re
import unicodedata
from collections import Counter
from datetime import datetime, timezone, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

CONTRACT = json.loads((Path(__file__).resolve().parents[1] / 'fourwall_contract.json').read_text(encoding='utf-8'))
SENSITIVE = re.compile(r'password|passwd|contrase[nñ]a|cookie|authorization|access.?token|refresh.?token|service.?role|api.?key', re.I)


class ValidationError(ValueError):
    def __init__(self, code):
        super().__init__(code)
        self.code = code


def tidy(value):
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFC', str(value)).strip())


def number(value, required=False):
    text = tidy(value).replace(',', '')
    if not text and not required:
        return ''
    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', text):
        raise ValidationError('INVALID_NUMBER')
    try:
        parsed = Decimal(text)
        exponent = int(text.lower().split('e')[1]) if 'e' in text.lower() else 0
        if not parsed.is_finite() or abs(parsed) > Decimal('1e12') or abs(exponent) > 20:
            raise ValidationError('INVALID_NUMBER')
        fixed = format(parsed, 'f')
        if '.' in fixed:
            fixed = fixed.rstrip('0').rstrip('.')
        return '0' if parsed == 0 else fixed.lstrip('+')
    except InvalidOperation as exc:
        raise ValidationError('INVALID_NUMBER') from exc


def date(value, date_order=None):
    if value is None or tidy(value) == '':
        return ''
    if isinstance(value, (int, float)):
        value = datetime(1899, 12, 30, tzinfo=timezone.utc) + timedelta(milliseconds=round(value * 86400000))
    if isinstance(value, datetime):
        result = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    else:
        text = tidy(value)
        local = re.fullmatch(r'(\d{1,2})[/-](\d{1,2})[/-](\d{4})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?', text)
        if local:
            a, b, year, hour, minute, second = local.groups()
            order = date_order or ('DMY' if int(a) > 12 else 'MDY' if int(b) > 12 else None)
            if order not in ('DMY', 'MDY'):
                raise ValidationError('AMBIGUOUS_DATE')
            try:
                result = datetime(int(year), int(b if order == 'DMY' else a), int(a if order == 'DMY' else b), int(hour or 0), int(minute or 0), int(second or 0), tzinfo=timezone.utc)
            except ValueError as exc:
                raise ValidationError('INVALID_DATE') from exc
        else:
            if not re.match(r'^\d{4}-\d{2}-\d{2}(?:$|[ T])', text):
                raise ValidationError('INVALID_DATE')
            try:
                result = datetime.fromisoformat(text.replace('Z', '+00:00'))
                result = result.replace(tzinfo=timezone.utc) if result.tzinfo is None else result.astimezone(timezone.utc)
            except ValueError as exc:
                raise ValidationError('INVALID_DATE') from exc
    return result.isoformat(timespec='milliseconds').replace('.000+00:00', 'Z').replace('+00:00', 'Z')


def hash_values(prefix, values):
    packed = prefix + ''.join(f'{len(str(v).encode("utf-8"))}:{v}' for v in values)
    return hashlib.sha256(packed.encode('utf-8')).hexdigest()


def normalize_record(row, date_order=None):
    if not isinstance(row, dict):
        raise ValidationError('INVALID_ROW')
    if any(SENSITIVE.search(str(key)) for key in row):
        raise ValidationError('UNSAFE_COLUMN')
    canonical, columns = [], {}
    for field in CONTRACT['fields']:
        present = [name for name in field['aliases'] if row.get(name) is not None]
        if len({tidy(row[name]) for name in present}) > 1:
            raise ValidationError('AMBIGUOUS_COLUMN')
        column = present[0] if present else None
        value = row.get(column)
        if field['key'] in ('ticket', 'serial', 'part', 'part_original') and isinstance(value, (int, float)):
            if not float(value).is_integer() or abs(value) >= 1e15:
                raise ValidationError('NUMERIC_ID_UNSAFE')
            value = str(int(value))
        text = number(value, field.get('required', False)) if field['kind'] == 'number' else date(value, date_order) if field['kind'] == 'date' else tidy(value)
        if field['kind'] == 'upper':
            text = text.upper()
        if field.get('required') and not text:
            raise ValidationError('REQUIRED_FIELD')
        if len(text) > 512:
            raise ValidationError('FIELD_TOO_LONG')
        canonical.append(text)
        if column:
            columns[field['key']] = column
    if not canonical[0]:
        raise ValidationError('MISSING_TICKET')
    raw = {str(k): v.isoformat() if isinstance(v, datetime) else v for k, v in row.items() if not str(k).startswith('__')}
    if any(isinstance(v, (dict, list, tuple)) for v in raw.values()):
        raise ValidationError('INVALID_ROW')
    if len(json.dumps(raw, ensure_ascii=False).encode('utf-8')) > 16384:
        raise ValidationError('ROW_TOO_LARGE')
    return dict(canonical=canonical, row_hash=hash_values('4wall-row-v1|', canonical), source_record_id=canonical[0], numero_parte=canonical[6], cantidad=float(canonical[7]), area_escaneo=canonical[1], raw_record=raw, source_columns={'part_number': columns['part'], 'quantity': columns['quantity'], 'area': columns['area']})


def prepare_snapshot(rows, *, complete=True, date_order=None, previous_count=0, min_row_ratio=0.1, max_rows=100000, max_quantity=1e9, export_complete=True):
    if not rows or len(rows) > max_rows:
        raise ValidationError('INVALID_ROW_COUNT')
    if complete and (not export_complete or previous_count and len(rows) < previous_count * min_row_ratio):
        raise ValidationError('INCOMPLETE_EXPORT')
    records = [normalize_record(row, date_order) for row in rows]
    counts = Counter(row['source_record_id'] for row in records)
    identities = set()
    for row in records:
        if not math.isfinite(row['cantidad']) or abs(row['cantidad']) > max_quantity:
            raise ValidationError('QUANTITY_LIMIT')
        collision = counts[row['source_record_id']] > 1
        if collision and (not row['canonical'][3] or not row['canonical'][12]):
            raise ValidationError('AMBIGUOUS_IDENTITY')
        row['identity_mode'] = 'composite' if collision else 'ticket'
        row['composite_key'] = hash_values('4wall-composite-v1|', [row['canonical'][0], row['canonical'][6], row['canonical'][12], row['canonical'][3]])
        row['source_identity'] = row['composite_key'] if collision else hash_values('4wall-ticket-v1|', [row['source_record_id']])
        if row['source_identity'] in identities:
            raise ValidationError('AMBIGUOUS_IDENTITY')
        identities.add(row['source_identity'])
    return dict(records=records, snapshot_hash=hash_values('4wall-snapshot-v1|', sorted(row['source_identity'] + ':' + row['row_hash'] for row in records)), rows_seen=len(records), complete=complete)


def build_batches(snapshot, current_state=(), batch_size=250):
    hashes = {row.get('row_hash') for row in current_state}
    records = [dict(row, position=i, raw_record=None if row['row_hash'] in hashes else row['raw_record']) for i, row in enumerate(snapshot['records'])]
    return [records[i:i + batch_size] for i in range(0, len(records), batch_size)]


def read_excel(path):
    """Read the export without integer-casting quantities or rounding text ticket identifiers."""
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        candidates = []
        for sheet in workbook:
            for header_index, cells in enumerate(sheet.iter_rows(min_row=1, max_row=50), start=1):
                headers = [tidy(c.value) for c in cells]
                if all(any(a in headers for a in f['aliases']) for f in CONTRACT['fields'] if f.get('required')) and any(a in headers for a in CONTRACT['fields'][0]['aliases']):
                    candidates.append((sheet, headers, header_index))
                    break
        if len(candidates) != 1:
            raise ValidationError('UNRECOGNIZED_SHEET')
        sheet, headers, header_index = candidates[0]
        if len(set(headers)) != len(headers) or any(not h for h in headers):
            raise ValidationError('AMBIGUOUS_COLUMN')
        rows = []
        for cells in sheet.iter_rows(min_row=header_index + 1):
            if all(cell.value is None for cell in cells):
                continue
            row = {}
            for header, cell in zip(headers, cells):
                if cell.data_type == 'f':
                    raise ValidationError('FORMULA_EXPORT')
                value = cell.value
                if header in ('Ticket/FIFO', 'serial', 'Serial', 'Número Parte QAD', 'Numero Parte QAD', 'Numero de parte', 'Número de parte', 'numero_parte') and isinstance(value, (int, float)):
                    if abs(value) >= 1e15 or not float(value).is_integer():
                        raise ValidationError('NUMERIC_ID_UNSAFE')
                    value = str(int(value))
                    if re.fullmatch('0+', cell.number_format):
                        value = value.zfill(len(cell.number_format))
                row[header] = value
            rows.append(row)
            if len(rows) > 250000:
                raise ValidationError('INVALID_ROW_COUNT')
        return rows
    finally:
        workbook.close()
