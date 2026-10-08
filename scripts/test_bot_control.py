"""Synthetic runner regressions: no network, real exports, credentials or listening ports."""
import io
import json
import tempfile
import unittest
import urllib.error
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from runner.sync import prepare_snapshot, read_excel, ValidationError
from runner.client import SupabaseRunnerClient, RemoteError
from runner.source import SourceError
from runner.engine import RunnerEngine

ROW = {'Ticket/FIFO': '000001', 'Número Parte QAD': 'PN-TEST', 'AreaName': 'AREA', 'Quantity': 1.25, 'Fecha agregado': '2026-10-08T08:00:00Z', 'Escaneador': 'SCANNER'}

class NoRealtime:
    def start(self): pass
    def close(self): pass

class Source:
    def __init__(self, fail=False, bad_login=False):
        self.fail, self.bad_login, self.calls, self.closed = fail, bad_login, 0, False
    def login(self):
        if self.bad_login: raise SourceError('LOGIN_ERROR', 'login')
    def extract(self):
        self.calls += 1
        if self.fail: raise SourceError('DOWNLOAD_ERROR', 'download')
        return {'rows': [ROW], 'extracted_at': '2026-10-08T08:00:00Z', 'export_complete': True}
    def close(self): self.closed = True

class Client:
    def __init__(self, lost_ack=False, paused=False):
        self.operations, self.receipt, self.attempt = [], {'status': 'RUNNING'}, 0
        self.lost_ack, self.paused, self.controls, self.engine = lost_ack, paused, 0, None
    def read_manifest(self): return 0, []
    def api(self, op, **p):
        self.operations.append(op)
        if op in ('start', 'control'):
            if op == 'control':
                self.controls += 1
                if self.engine and self.controls >= 3: self.engine.request_close()
            return {'runner': {'enabled': True, 'paused': self.paused, 'next_run_at': '2100-01-01T00:00:00Z'}, 'source': {'generation': 0, 'row_count': 1, 'min_row_ratio': .1, 'max_rows': 100000, 'max_quantity': 1e9}, 'pending': []}
        if op == 'begin': return {'lease_token': 'synthetic-lease'}
        if op == 'publish':
            self.receipt = {'status': 'SUCCESS', 'id': p['run_id'], 'rows_seen': 1, 'rows_inserted': 1}
            if self.lost_ack: raise RemoteError()
            return self.receipt
        if op == 'receipt': return self.receipt
        if op in ('retry', 'fail'):
            self.attempt += 1
            self.receipt = {'status': 'FAILED' if op == 'fail' else 'RETRYING', 'next_attempt_at': '2000-01-01T00:00:00Z'}
            return self.receipt
        return {}

class Tests(unittest.TestCase):
    def test_identity_and_fractional_parser(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic.xlsx'
            book = Workbook(); sheet = book.active
            sheet.append(list(ROW)); sheet.append(list(ROW.values())); book.save(path); book.close()
            parsed = read_excel(path)
            self.assertEqual(parsed[0]['Ticket/FIFO'], '000001')
            self.assertEqual(parsed[0]['Quantity'], 1.25)
            self.assertEqual(prepare_snapshot(parsed)['snapshot_hash'], prepare_snapshot([ROW])['snapshot_hash'])
            book = Workbook(); sheet = book.active; sheet.append(list(ROW)); sheet.append(list(ROW.values())); sheet.cell(2, 4, '=1+2');book.save(path);book.close()
            with self.assertRaisesRegex(ValidationError, 'FORMULA_EXPORT'): read_excel(path)
    def test_collisions_are_not_deduplicated(self):
        rows = [dict(ROW, Escaneador=f'S-{i}', **{'Número Parte QAD': f'PN-{i%2}'}) for i in range(6)]
        snapshot = prepare_snapshot(rows)
        self.assertEqual(len({r['source_identity'] for r in snapshot['records']}), 6)
        self.assertTrue(all(r['identity_mode'] == 'composite' for r in snapshot['records']))
        with self.assertRaisesRegex(ValidationError, 'AMBIGUOUS_IDENTITY'): prepare_snapshot([ROW, ROW])
        with self.assertRaises(ValidationError): prepare_snapshot([dict(ROW, Quantity='invalid')])
    def test_finite_retries_and_no_parallel(self):
        source, client = Source(fail=True), Client(); engine = RunnerEngine(client, source, wakeup=NoRealtime())
        self.assertEqual(engine.execute('SCHEDULE', {})['status'], 'FAILED')
        self.assertEqual(source.calls, 4); self.assertEqual(client.operations.count('retry'), 3)
        self.assertFalse(engine.active)
        engine.active = True
        with self.assertRaisesRegex(RuntimeError, 'PARALLEL'): engine.execute('SCHEDULE', {})
    def test_lost_ack_returns_committed_receipt(self):
        source, client = Source(), Client(lost_ack=True); engine = RunnerEngine(client, source, wakeup=NoRealtime())
        self.assertEqual(engine.execute('SCHEDULE', {})['status'], 'SUCCESS')
        self.assertEqual(source.calls, 1); self.assertNotIn('fail', client.operations); self.assertNotIn('retry', client.operations)
    def test_invalid_login_does_not_extract(self):
        source, client = Source(bad_login=True), Client(); engine = RunnerEngine(client, source, wakeup=NoRealtime()); engine.run()
        self.assertNotIn('begin', client.operations); self.assertEqual(source.calls, 0); self.assertTrue(source.closed)
    def test_close_during_login_does_not_start_new_cut(self):
        source, client = Source(), Client(); engine = RunnerEngine(client, source, wakeup=NoRealtime())
        source.login = engine.request_close
        engine.run()
        self.assertEqual(source.calls, 0); self.assertTrue(source.closed)
    def test_machine_claim_precedes_edge_login(self):
        source, client = Source(), Client(); engine = RunnerEngine(client, source, wakeup=NoRealtime())
        original = client.api
        def reject_start(op, **p):
            if op == 'start': raise RemoteError('RUNNER_ALREADY_ONLINE')
            return original(op, **p)
        client.api = reject_start
        source.login = lambda: self.fail('Second machine must not launch Edge')
        engine.run(); self.assertEqual(source.calls, 0); self.assertTrue(source.closed)
    def test_polling_survives_unavailable_realtime_and_pause(self):
        for paused in (False, True):
            source, client = Source(), Client(paused=paused); engine = RunnerEngine(client, source, wakeup=NoRealtime()); client.engine = engine
            class ImmediateWake:
                def wait(self, seconds): self.seconds = seconds
                def set(self): pass
                def clear(self): pass
            engine.wake = ImmediateWake(); engine.run()
            self.assertGreaterEqual(client.controls, 3)
            self.assertEqual(engine.wake.seconds, 5)
            self.assertEqual(source.calls, 0 if paused else 1)
            self.assertTrue(source.closed)
    def test_outbound_only_transport_and_safe_errors(self):
        with self.assertRaises(ValueError): SupabaseRunnerClient('http://localhost:8000', 'public', 'runner-test')
        with self.assertRaises(ValueError): SupabaseRunnerClient('https://uukhwkywmnarcfruerpp.supabase.co', 'sb_secret_not-a-key', 'runner-test')
        def opener(request, timeout):
            self.assertEqual(timeout, 45); self.assertTrue(request.full_url.startswith('https://uukhwkywmnarcfruerpp.supabase.co/'))
            raise urllib.error.HTTPError(request.full_url, 403, 'untrusted', {}, io.BytesIO(b'{"message":"UNTRUSTED_RESPONSE_BODY"}'))
        client = SupabaseRunnerClient('https://uukhwkywmnarcfruerpp.supabase.co', 'sb_publishable_synthetic_test', 'runner-test', opener=opener)
        client.access_token, client.expires_at = 'memory-only-test', 10**12
        with self.assertRaisesRegex(RemoteError, '^REMOTE_ERROR$'): client.api('control')
    def test_legacy_entrypoint_has_no_listener(self):
        import bot_control_server
        self.assertFalse(hasattr(bot_control_server, 'ThreadingHTTPServer'))
        import bot_extractor
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'config.json';path.write_text('{"password":"FORBIDDEN_CONFIG_FIELD"}')
            with self.assertRaises(ValueError): bot_extractor.load_config(path)

if __name__ == '__main__': unittest.main(verbosity=2)
